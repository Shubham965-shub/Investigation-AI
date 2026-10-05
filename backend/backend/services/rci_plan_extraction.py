"""Extracts Task Critique's task list from our real RCI Plan Word template's structure — the reverse of rci_plan_export.py's cell-write navigation.
Built here rather than calling ds's POST /critique/extract, whose parser expects a differently-structured document and fails against our template.
"""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional

import docx


def _clean(text: Optional[str]) -> Optional[str]:
    """Empty string and the literal "Unassigned"/"N/A" placeholders all mean "nothing set" — normalize all to None."""
    text = (text or "").strip()
    if not text or text in ("Unassigned", "N/A"):
        return None
    return text


_SUBTASK_NUMBER_RE = re.compile(r"^\d+\.\d+\s*")


def extract_task_sections(docx_bytes: bytes) -> List[Dict[str, Any]]:
    """Returns one dict per task row, in document order:
    {"title": str, "correlation": str | None, "tasks": list[str],
    "due_date": str | None, "assignee": str | None}.
    """
    doc = docx.Document(io.BytesIO(docx_bytes))
    outer = doc.tables[0]
    # Locate the "2.0 Investigation tasks" header row by its section-number label rather than a
    # hardcoded row index — a fixed index broke every already-exported docx from before the "1.5
    # Pre-requisite of Investigation Plan" checklist was inserted (2026-10-01), which shifted this
    # row by +2 and crashed Task Critique with an IndexError for ~90% of existing records. The
    # "2.0" cell itself is stable across both the old and new template layouts.
    header_row_index = next(
        (i for i, row in enumerate(outer.rows) if (row.cells[0].text or "").strip() == "2.0"),
        None,
    )
    if header_row_index is None:
        raise ValueError("Could not locate the \"2.0 Investigation tasks\" section in this RCI Plan export")
    tasks_table = outer.rows[header_row_index + 1].cells[1].tables[0]

    sections: List[Dict[str, Any]] = []
    for row in tasks_table.rows[1:]:  # row 0 is the header
        objective = (row.cells[1].text or "").strip()
        if not objective:
            continue

        # Mirrors rci_plan_export.py's f"{title}\n{correlation}" write — split back apart the same way.
        title, _, correlation = objective.partition("\n")

        # Mirrors rci_plan_export.py's "<task>.<subtask> " numbering — strip it back off. Zero checked subtasks exports as literal "N/A", not blank.
        task_lines = [
            _SUBTASK_NUMBER_RE.sub("", line.strip())
            for line in (row.cells[2].text or "").split("\n")
            if line.strip() and line.strip() != "N/A"
        ]

        sections.append(
            {
                "title": title.strip(),
                "correlation": correlation.strip() or None,
                "tasks": task_lines,
                "due_date": _clean(row.cells[4].text),
                "assignee": _clean(row.cells[3].text),
            }
        )

    return sections