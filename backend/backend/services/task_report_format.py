"""Validates an uploaded task report against the real Task Report Template's structure (before sending to ds for critique) — a structural label check, not a content/quality check.
"Title of the task:" is deliberately excluded from REQUIRED_MARKERS: once filled in, the title becomes a plain heading with no literal prefix, so requiring it would reject well-formed real reports.
"""
from __future__ import annotations

import io
from typing import List

import docx
from docx.oxml.ns import qn

REQUIRED_MARKERS = [
    "problem statement",
    "investigation tasks",
    "objective",
    "findings",
    "inference",
]


def missing_task_report_markers(file_bytes: bytes) -> List[str]:
    """Returns REQUIRED_MARKERS missing from the document body (case-insensitive); empty list means it checks out.
    Walks every <w:t> node via the raw XML tree rather than doc.paragraphs/doc.tables, which miss text nested inside a table-within-a-table or a content control (<w:sdt>) block."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception:
        return list(REQUIRED_MARKERS)

    body_text = "".join(node.text or "" for node in doc.element.body.iter(qn("w:t")))
    body_text_lower = body_text.lower()

    return [marker for marker in REQUIRED_MARKERS if marker not in body_text_lower]


def is_valid_task_report_format(file_bytes: bytes) -> bool:
    return not missing_task_report_markers(file_bytes)