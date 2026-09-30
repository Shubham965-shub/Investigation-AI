import io

import docx

from backend.services.task_report_format import (
    REQUIRED_MARKERS,
    is_valid_task_report_format,
    missing_task_report_markers,
)

_ALL_LABELS = ["Problem statement:", "Investigation tasks", "Objective:", "Findings:", "Inference:"]


def _docx_bytes(paragraphs=(), table_cells=()):
    doc = docx.Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_cells:
        table = doc.add_table(rows=1, cols=len(table_cells))
        for cell, text in zip(table.rows[0].cells, table_cells):
            cell.text = text
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_all_labels_present_passes():
    assert missing_task_report_markers(_docx_bytes(_ALL_LABELS)) == []
    assert is_valid_task_report_format(_docx_bytes(_ALL_LABELS))


def test_labels_are_case_insensitive():
    assert missing_task_report_markers(_docx_bytes([label.upper() for label in _ALL_LABELS])) == []


def test_missing_label_is_reported():
    # The 510613 case: the Inference heading was overwritten with "Findings:".
    labels = [label for label in _ALL_LABELS if label != "Inference:"]
    assert missing_task_report_markers(_docx_bytes(labels)) == ["inference"]
    assert not is_valid_task_report_format(_docx_bytes(labels))


def test_labels_inside_tables_are_found():
    assert missing_task_report_markers(_docx_bytes(table_cells=_ALL_LABELS)) == []


def test_label_split_across_runs_is_found():
    doc = docx.Document()
    para = doc.add_paragraph()
    para.add_run("Infer")
    para.add_run("ence:")
    for label in _ALL_LABELS[:-1]:
        doc.add_paragraph(label)
    buffer = io.BytesIO()
    doc.save(buffer)
    assert missing_task_report_markers(buffer.getvalue()) == []


def test_unreadable_file_reports_every_marker():
    assert missing_task_report_markers(b"not a docx") == list(REQUIRED_MARKERS)
