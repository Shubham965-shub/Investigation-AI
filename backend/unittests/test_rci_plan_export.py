import io

import docx

from backend.schemas.rci_plan import RciPrerequisiteChecklist, RciSectionItem, RciTaskItem
from backend.services.rci_plan_export import MAX_TEMPLATE_SECTIONS, _format_ddmmyyyy, build_rci_plan_docx
from backend.services.rci_plan_extraction import extract_task_sections


def _section(title, tasks=("Check the thing",), **kwargs):
    return RciSectionItem(title=title, tasks=[RciTaskItem(description=t) for t in tasks], **kwargs)


def _build(sections):
    return build_rci_plan_docx("DEV-1", {"description": "Tablet hardness out of limit"}, sections, RciPrerequisiteChecklist())


def test_format_ddmmyyyy():
    assert _format_ddmmyyyy("2026-09-30") == "30/09/2026"
    assert _format_ddmmyyyy(None) == ""
    assert _format_ddmmyyyy("") == ""
    assert _format_ddmmyyyy("not-a-date-value") == "not-a-date-value"


def test_build_then_extract_round_trips_sections():
    sections = [
        _section("Machine", ["Review logbook", "Check calibration"], correlation="Linked to punch wear", assignee="Asha", due_date="2026-10-05"),
        _section("Method", ["Review BMR"]),
    ]
    docx_bytes, truncated, owners_truncated = _build(sections)
    assert (truncated, owners_truncated) == (0, 0)

    extracted = extract_task_sections(docx_bytes)
    assert [s["title"] for s in extracted] == ["Machine", "Method"]
    assert extracted[0]["correlation"] == "Linked to punch wear"
    assert extracted[0]["tasks"] == ["Review logbook", "Check calibration"]
    assert extracted[0]["assignee"] == "Asha"
    assert extracted[0]["due_date"] == "05/10/2026"
    # Blank assignee/due date export as placeholders, which extraction normalizes back to None.
    assert extracted[1]["assignee"] is None
    assert extracted[1]["due_date"] is None


def test_unchecked_sections_and_tasks_are_excluded():
    sections = [
        _section("Kept", tasks=[]),
        _section("Dropped", is_checked=False),
    ]
    sections[0].tasks = [RciTaskItem(description="Keep me"), RciTaskItem(description="Skip me", is_checked=False)]
    extracted = extract_task_sections(_build(sections)[0])
    assert [s["title"] for s in extracted] == ["Kept"]
    assert extracted[0]["tasks"] == ["Keep me"]


def test_sections_beyond_template_slots_are_truncated():
    sections = [_section(f"Task {i}") for i in range(MAX_TEMPLATE_SECTIONS + 2)]
    docx_bytes, truncated, _ = _build(sections)
    assert truncated == 2
    assert len(extract_task_sections(docx_bytes)) == MAX_TEMPLATE_SECTIONS


def test_more_than_four_owners_adds_sign_off_rows():
    sections = [_section(f"Task {i}", assignee=f"Owner {i}") for i in range(6)]
    docx_bytes, _, owners_truncated = _build(sections)
    assert owners_truncated == 0
    text = "\n".join(cell.text for table in docx.Document(io.BytesIO(docx_bytes)).tables for row in table.rows for cell in row.cells)
    assert "Owner 5" in text
    assert "Task Owner 5" in text
