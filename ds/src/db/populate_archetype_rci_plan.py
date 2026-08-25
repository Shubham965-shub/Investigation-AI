import os
import re
import sys
from pathlib import Path

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from config.settings import settings
from archetype_models import Base, ArchetypeType, Archetype, RciPlan, RciPlanSection, RciPlanTask

EXCEL_PATH = os.path.join(
    os.path.dirname(__file__),
    "RCI_Plan_Library_6M_grammar_checked.xlsx",
)

engine = create_engine(settings.DATABASE_URL, connect_args={"sslmode": "require"})
Session = sessionmaker(bind=engine)
session = Session()


def clean_text(text: str) -> str:
    """Normalize non-breaking spaces/whitespace and strip a leading bullet marker."""
    text = str(text).replace("\xa0", " ")
    text = re.sub(r"^[•\-\*]\s*", "", text.strip())
    return re.sub(r"\s+", " ", text).strip()


def split_title(raw: str) -> tuple[str, str | None]:
    """Split a 'Task(s) & objective' cell into (title, correlation).

    Cells are authored as 'Title (explanation)' or 'Title\\n(explanation)'.
    The title itself can contain its own parenthetical abbreviation
    (e.g. 'Material Attributes (CMA) & Charging\\n(Rule out ...)'), so when a
    newline is present we split on the *last* line rather than the first
    '(' - the explanation is always the final line. Only single-line cells
    fall back to splitting on the first '('.
    """
    text = str(raw).replace("\xa0", " ").strip()
    if "\n" in text:
        head, _, tail = text.rpartition("\n")
        title = re.sub(r"\s+", " ", head).strip()
        tail = tail.strip()
        correlation = tail[1:-1] if tail.startswith("(") and tail.endswith(")") else tail
        correlation = re.sub(r"\s+", " ", correlation).strip()
        return title, correlation or None
    idx = text.find("(")
    if idx == -1:
        return re.sub(r"\s+", " ", text).strip(), None
    title = re.sub(r"\s+", " ", text[:idx]).strip()
    correlation = text[idx:]
    if correlation.startswith("(") and correlation.endswith(")"):
        correlation = correlation[1:-1]
    else:
        correlation = correlation[1:]
    correlation = re.sub(r"\s+", " ", correlation).strip()
    return title, correlation or None


def load_rci_plans_from_excel(path: str) -> dict[str, list[list]]:
    """Return {archetype_name: [[section_title, correlation, six_m_bucket, [task, ...]], ...]}.

    Sheet columns are Sr.No. | 6M Bucket | Task(s) & objective | Task details.
    A new Sr.No. marks a new section; some sections author their explanation
    on its own follow-up row (objective cell is just '(explanation)') instead
    of appending it to the title cell with a newline - both forms are handled.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    plans: dict[str, list[list]] = {}
    for sheet_name in wb.sheetnames:
        if sheet_name == "Index":
            continue
        ws = wb[sheet_name]
        sections: list[list] = []
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
            _, six_m, objective, detail = (row + (None, None, None, None))[:4]
            if objective:
                text = str(objective).strip()
                if sections and sections[-1][1] is None and text.startswith("(") and text.endswith(")"):
                    sections[-1][1] = clean_text(text[1:-1])
                else:
                    title, correlation = split_title(text)
                    bucket = str(six_m).strip() if six_m else None
                    sections.append([title, correlation, bucket, []])
            if detail and sections:
                sections[-1][3].append(clean_text(detail))
        plans[sheet_name.strip()] = sections
    return plans


# 1. Ensure archetype_type exists
archetype_type = session.query(ArchetypeType).filter_by(name="RCI Plan").first()
if not archetype_type:
    archetype_type = ArchetypeType(name="RCI Plan")
    session.add(archetype_type)
    session.commit()

# 2. Delete existing archetypes (and cascaded plans/sections/tasks) for this type
existing = session.query(Archetype).filter_by(archetype_type_id=archetype_type.id).all()
for arch in existing:
    plan = session.query(RciPlan).filter_by(archetype_id=arch.id).first()
    if plan:
        session.delete(plan)  # cascades to sections -> tasks
    session.delete(arch)
session.commit()
print(f"Removed {len(existing)} existing archetype(s) for 'RCI Plan'.")

# 3. Load new data from Excel
plans = load_rci_plans_from_excel(EXCEL_PATH)

# 4. Insert new archetypes with their RCI plan / sections / tasks
for name, sections in plans.items():
    archetype = Archetype(name=name, definition=None, archetype_type_id=archetype_type.id)
    session.add(archetype)
    session.flush()

    rci_plan = RciPlan(archetype_id=archetype.id)
    session.add(rci_plan)
    session.flush()

    for title, correlation, six_m_bucket, tasks in sections:
        section = RciPlanSection(
            rci_plan_id=rci_plan.id,
            title=title,
            correlation=correlation,
            six_m_bucket=six_m_bucket,
        )
        session.add(section)
        session.flush()
        for task_desc in tasks:
            session.add(RciPlanTask(rci_plan_section_id=section.id, description=task_desc))

session.commit()
total_sections = sum(len(s) for s in plans.values())
total_tasks = sum(len(t) for _, _, _, t in (sec for s in plans.values() for sec in s))
print(
    f"Populated {len(plans)} archetype(s) with {total_sections} section(s) and "
    f"{total_tasks} task(s) from '{os.path.basename(EXCEL_PATH)}'."
)
