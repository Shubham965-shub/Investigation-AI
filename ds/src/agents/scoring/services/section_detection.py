"""
Detect which of the four scoreable sections a report contains and split them out.

Accepts .docx and .pdf reports.

Primary path: upload the file and let the model read the native structure
(reuses LLMClient.upload_file / get_structured_response_from_file; the Responses
API reads both DOCX and PDF). If the upload path fails (e.g. no /v1/files access),
fall back to text extracted locally — python-docx for .docx, pypdf for .pdf — plus
a chat call.
"""

from __future__ import annotations

import asyncio
import logging
from pathlib import Path
from typing import Dict

from src.agents.scoring.api.schemas import SectionDetectionLLMOutput
from src.agents.scoring.support import SCORING_SYSTEM_PREAMBLE
from src.llm.client import LLMClient
from src.prompt_registry.service import PromptRegistry

logger = logging.getLogger(__name__)

_SECTION_FIELDS = {
    "task_report": "task_report_text",
    "rc": "rc_text",
    "impact": "impact_text",
    "capa": "capa_text",
}

_EVENT_TYPE_CANON = {
    "deviation": "Deviation",
    "oos": "OOS",
    "out of specification": "OOS",
    "oot": "OOT",
    "out of trend": "OOT",
    "market complaint": "Market Complaint",
    "complaint": "Market Complaint",
}


def normalise_event_type(value: str | None) -> str | None:
    if not value:
        return None
    return _EVENT_TYPE_CANON.get(value.strip().lower(), value.strip())


def present_sections(detection: SectionDetectionLLMOutput) -> Dict[str, str]:
    """Return {section_key: text} for every section that has non-blank content."""
    out: Dict[str, str] = {}
    for key, field in _SECTION_FIELDS.items():
        text = (getattr(detection, field) or "").strip()
        if text:
            out[key] = text
    return out


def docx_to_text(file_path: Path) -> str:
    """Flatten a .docx to plain text (paragraphs + tables) using python-docx.

    Reliable local fallback when the LLM file-upload path is unavailable — the
    reports are table-heavy, so table cells are included row by row.
    """
    from docx import Document  # local import: only needed on the fallback path

    doc = Document(str(file_path))
    parts: list[str] = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


def pdf_to_text(file_path: Path) -> str:
    """Extract plain text from a .pdf using pypdf (fallback path only)."""
    from pypdf import PdfReader  # local import: only needed on the fallback path

    reader = PdfReader(str(file_path))
    parts: list[str] = []
    for page in reader.pages:
        t = (page.extract_text() or "").strip()
        if t:
            parts.append(t)
    return "\n".join(parts)


def file_to_text(file_path: Path) -> str:
    """Extract text from a supported report file (.docx or .pdf)."""
    if file_path.suffix.lower() == ".pdf":
        return pdf_to_text(file_path)
    return docx_to_text(file_path)


async def detect_and_split(
    file_path: Path, llm: LLMClient, registry: PromptRegistry
) -> SectionDetectionLLMOutput:
    prompt = registry.get("scoring/detect_sections")

    # Primary: native file upload.
    try:
        file_id = await llm.upload_file(file_path)
        try:
            result: SectionDetectionLLMOutput = await llm.get_structured_response_from_file(
                file_id=file_id,
                user_prompt=prompt,
                structure=SectionDetectionLLMOutput,
                system_prompt=SCORING_SYSTEM_PREAMBLE,
                temperature=0,
            )
        finally:
            await llm.delete_file(file_id)
        result.event_type = normalise_event_type(result.event_type)
        logger.info("detect_and_split: file-upload path → sections %s", list(present_sections(result)))
        return result
    except Exception as exc:  # noqa: BLE001 - fall back on any upload failure
        logger.warning("Section detection file-upload failed (%s); using local text fallback.", exc)

    # Fallback: extract text locally (docx via python-docx, pdf via pypdf), then a chat call.
    text = await asyncio.to_thread(file_to_text, file_path)
    if not text.strip():
        raise RuntimeError(f"Could not extract any text from {file_path} for section detection.")

    user_prompt = f"{prompt}\n\nDOCUMENT TEXT:\n{text}"
    result = await llm.get_structured_chat_response(
        user_prompt=user_prompt,
        structure=SectionDetectionLLMOutput,
        system_prompt=SCORING_SYSTEM_PREAMBLE,
    )
    result.event_type = normalise_event_type(result.event_type)
    logger.info("detect_and_split: local-text path → sections %s", list(present_sections(result)))
    return result
