from __future__ import annotations

import asyncio
import base64
import json
import logging
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from docx.oxml.ns import qn
from markitdown import MarkItDown

from src.agents.critique.api.services.relevance_validation import validate_document_relevance
from src.agents.critique.graph.schemas import (
    AllTaskCritiquesResult,
    ExtractionResult,
    ImageAnalysisResult,
    SectionTasks,
    TaskReportCritiqueResponse,
    _ExtractionHeader,
)
from src.agents.critique.graph.state import TaskReportCritiqueState
from src.llm.client import LLMClient
from src.prompt_registry.service import PromptRegistry
from src.utils.deps import get_db_pool, get_llm_client, get_prompt_registry

_markitdown = MarkItDown()

logger = logging.getLogger(__name__)

# Fixed leading phrase a regenerated recommendation uses when the previous attempt's accepted
# item is still not addressed by the new report (see critique_task.yaml v6). Also the marker
# critique_tasks' post-LLM cap looks for to prioritize these over brand-new recommendations.
UNADDRESSED_MARKER = "Still unaddressed from the previous review."

# ── Markitdown + inference-count extraction helpers ───────────────────────────

_INF_RE = re.compile(r'\*{0,2}\s*Inference\s*:?\*{0,2}', re.I)
# Match end-of-tasks section headings only — anchored to line start, not table rows.
# "Causal Factor" requires a trailing colon to avoid matching body text like
# "causal factors can be ruled out."
_END_RE = re.compile(
    r'^(?!\s*\|)\s*\*{0,2}\s*(?:Identified\s+Root\s+Cause|Causal\s+[Ff]actors?\s*:|CAPA\s+Proposal)',
    re.I | re.MULTILINE,
)
# Top-level numbered section: "1. Title", "# 1. Title", "**1. Title**" — single number only
_SECTION_RE = re.compile(
    r'^(?:#{1,3}\s+|\*{0,2}\s*)(\d+)\.(?!\d)\s+\S',
    re.MULTILINE,
)

_HEADER_PROMPT = """You are a pharmaceutical QA expert.
Extract the problem_statement and investigation objective from this task report content.

TEXT:
{context}

Return JSON with:
  problem_statement: full narrative under "Problem statement:" — copy verbatim, do not paraphrase.
  objective: framing text under "Investigation tasks:" (before first numbered task). Blank if absent.

Return ONLY valid JSON. No markdown fences."""

_FULL_DOC_PROMPT = """You are a pharmaceutical QA expert extracting investigation tasks from the attached task report document.

The document template for every task is:
─────────────────────────────────────────
<Title / bold label / numbered heading>
  Objective: <what this task aims to determine>     ← sometimes absent
  <findings body — text + tables>
Inference: <concluding inference for this task>     ← always present
─────────────────────────────────────────

This document has exactly {n_total} "Inference:" markers across all investigation sections.
Extract exactly {n_total} tasks — one per Inference:, in document order.
Task titles may be bold headings, numbered sub-headings, OR first-column table labels (e.g. material names).
Do NOT extract Identified Root Cause, Causal Factor, or CAPA Proposal sections.

Extract:
1. problem_statement — full narrative under "Problem statement:". Copy verbatim.
2. objective — framing text under "Investigation tasks:" before the first task. Blank if absent.
3. tasks — exactly {n_total} items:
   task_number, title, objective, findings (include table content verbatim), inference, section_labels.

Return ONLY valid JSON matching the schema. No markdown fences."""

_SECTION_PROMPT = """You are a pharmaceutical QA expert. Below is one investigation section from a pharmaceutical task report.

This section contains exactly {n_inf} "Inference:" marker(s).
Extract exactly {n_inf} task(s) — one per Inference: occurrence.

Key rules:
- Every Inference: ends exactly one task.
- Task titles may be bold headings, numbered sub-headings, OR bold text / first-column labels inside a \
table (e.g. material names like "Cortisone Acetate").
- If a title is embedded in a table row, use the first prominent label (bold or first cell) preceding \
that task's content as the title.
- objective: text after "Objective:" for this task. Blank if absent.
- findings: all body text and table rows between this task's Objective and its Inference. \
Include table content verbatim.
- inference: text after "Inference:" for this task only.
- section_labels: header text of every table or sub-heading inside this task.

SECTION HEADING: {heading}

SECTION CONTENT:
{content}

Return JSON: {{ "tasks": [ {{task_number, title, objective, findings, inference, section_labels}}, ... ] }}
task_number values start at {start_num}.
Return ONLY valid JSON. No markdown fences."""


def _task_body(markdown: str) -> str:
    """Return markdown up to the first Root Cause / CAPA Proposal section heading."""
    end = _END_RE.search(markdown)
    return markdown[: end.start()] if end else markdown


def _segment_sections(body: str) -> list[tuple[str, str]]:
    """
    Split body at top-level numbered headings (1., 2., 3. — not 1.1, 2.3).
    Returns (section_heading_text, section_content) pairs that contain Inference: markers.
    """
    matches = list(_SECTION_RE.finditer(body))
    if not matches:
        return [("", body)]
    segments = []
    for i, m in enumerate(matches):
        heading = body[m.start(): m.end()].strip().strip("*# ").strip()
        content_start = m.end()
        content_end = matches[i + 1].start() if i + 1 < len(matches) else len(body)
        content = body[content_start:content_end].strip()
        if _INF_RE.search(content):
            segments.append((heading, content))
    return segments


# ── Helpers ──────────────────────────────────────────────────────────────────

def _extract_images_from_docx(doc: Document) -> List[Dict[str, Any]]:
    """
    Extract all images with their surrounding context label.
    Returns list of {name, b64, media_type, context_label}.
    """
    # Build a map: image part name → bytes
    part_map: Dict[str, tuple[bytes, str]] = {}
    for part in doc.part.package.iter_parts():
        if "image" in part.content_type:
            part_map[part.partname] = (part.blob, part.content_type)

    images: List[Dict[str, Any]] = []
    seen_names: set[str] = set()

    def _collect_from_element(element, context_label: str) -> None:
        for inline in element.findall(".//" + qn("wp:inline"), element.nsmap):
            _add_image(inline, context_label)
        for anchor in element.findall(".//" + qn("wp:anchor"), element.nsmap):
            _add_image(anchor, context_label)

    def _add_image(container_elem, context_label: str) -> None:
        for blip in container_elem.findall(".//" + qn("a:blip"), container_elem.nsmap):
            r_embed = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            if not r_embed:
                continue
            # Resolve relationship to part name
            try:
                rel = doc.part.rels.get(r_embed)
                if rel is None:
                    continue
                part_name = rel.target_part.partname
            except Exception:
                continue
            if part_name in seen_names:
                continue
            seen_names.add(part_name)
            if part_name not in part_map:
                continue
            blob, content_type = part_map[part_name]
            images.append({
                "name": part_name.split("/")[-1],
                "b64": base64.b64encode(blob).decode(),
                "media_type": content_type,
                "context_label": context_label,
            })

    # Collect from paragraph-level images
    for para in doc.paragraphs:
        _collect_from_element(para._element, para.text.strip() or "General section")

    # Collect from tables — use the header cell text as context label
    for table in doc.tables:
        header_text = table.rows[0].cells[0].text.strip() if table.rows else "Table"
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _collect_from_element(para._element, header_text)

    return images


def _build_raw_text(doc: Document) -> tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    paragraphs = [
        {"idx": i, "style": p.style.name, "text": p.text.strip()}
        for i, p in enumerate(doc.paragraphs)
        if p.text.strip()
    ]
    tables = []
    for t_idx, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"table_idx": t_idx, "rows": rows})
    return paragraphs, tables


# ── Nodes ─────────────────────────────────────────────────────────────────────

async def parse_document(state: TaskReportCritiqueState) -> Dict[str, Any]:
    """Extract paragraphs, tables, and images from the uploaded docx."""
    doc = await asyncio.to_thread(Document, state.file_path)
    paragraphs, tables = await asyncio.to_thread(_build_raw_text, doc)
    image_context_map = await asyncio.to_thread(_extract_images_from_docx, doc)
    logger.info(
        "parse_document: %d paragraphs, %d tables, %d images",
        len(paragraphs), len(tables), len(image_context_map),
    )
    return {
        "raw_paragraphs": paragraphs,
        "raw_tables": tables,
        "image_context_map": image_context_map,
    }


async def validate_relevance(state: TaskReportCritiqueState) -> Dict[str, Any]:
    """Fail fast if the uploaded document doesn't genuinely pertain to the given
    problem statement, before running the expensive extraction/critique LLM calls."""
    llm: LLMClient = await get_llm_client()
    document_text = "\n".join(p["text"] for p in state.raw_paragraphs)
    await validate_document_relevance(
        llm,
        problem_statement=state.problem_statement,
        event_type=state.event_type,
        document_text=document_text,
        document_label="investigation task report",
        task_context=state.task_description,
    )
    return {}


async def _extract_via_file_upload(
    state: TaskReportCritiqueState,
    llm: LLMClient,
    n_total: int,
) -> ExtractionResult:
    """Upload DOCX; LLM reads native structure. Requires /v1/files access."""
    file_id = await llm.upload_file(Path(state.file_path))
    try:
        result: ExtractionResult = await llm.get_structured_response_from_file(
            file_id=file_id,
            user_prompt=_FULL_DOC_PROMPT.format(n_total=n_total),
            structure=ExtractionResult,
        )
    finally:
        await llm.delete_file(file_id)
    return result


async def _extract_via_markitdown(
    state: TaskReportCritiqueState,
    llm: LLMClient,
    body: str,
    segments: list,
) -> ExtractionResult:
    """
    Fallback: markitdown text → per-section LLM extraction.
    Uses everything before the first Inference: marker as header context so
    the problem statement section is included even if it sits under a heading.
    """
    # Header context: everything before the first inference marker
    # This captures "1. Problem Statement" section content even though it's under a heading
    first_inf = _INF_RE.search(body)
    header_context = body[: first_inf.start()].strip() if first_inf else body[:4000]

    header: _ExtractionHeader = await llm.get_structured_chat_response(
        user_prompt=_HEADER_PROMPT.format(context=header_context),
        structure=_ExtractionHeader,
    )

    all_tasks = []
    task_counter = 1
    for heading, content in segments:
        n_inf = len(_INF_RE.findall(content))
        if n_inf == 0:
            continue
        prompt = _SECTION_PROMPT.format(
            n_inf=n_inf,
            heading=heading,
            content=content,
            start_num=task_counter,
        )
        sec: SectionTasks = await llm.get_structured_chat_response(
            user_prompt=prompt,
            structure=SectionTasks,
        )
        all_tasks.extend(sec.tasks)
        task_counter += len(sec.tasks)

    return ExtractionResult(
        problem_statement=header.problem_statement,
        objective=header.objective,
        tasks=all_tasks,
    )


async def extract_tasks(state: TaskReportCritiqueState) -> Dict[str, Any]:
    """
    1. markitdown → segment sections + count inferences (always runs; quick/local).
    2. Try file upload → Responses API (LLM reads native DOCX; best quality).
    3. On upload failure (network block etc.) fall back to per-section markitdown extraction.
    """
    llm: LLMClient = await get_llm_client()

    markdown = await asyncio.to_thread(
        lambda: _markitdown.convert(state.file_path).text_content
    )
    body = _task_body(markdown)
    segments = _segment_sections(body)
    n_total = sum(len(_INF_RE.findall(c)) for _, c in segments)

    try:
        result = await _extract_via_file_upload(state, llm, n_total)
        logger.info("extract_tasks: file upload path → %d tasks", len(result.tasks))
    except Exception as exc:
        logger.warning("File upload failed (%s); using markitdown fallback.", exc)
        result = await _extract_via_markitdown(state, llm, body, segments)
        logger.info("extract_tasks: markitdown path → %d tasks", len(result.tasks))

    return {
        "report_metadata": {
            "problem_statement": result.problem_statement,
            "objective": result.objective,
        },
        "extracted_tasks": [t.model_dump() for t in result.tasks],
    }


def _parse_json_column(raw: Any) -> Any:
    """asyncpg returns JSONB columns as text unless a codec is registered (none is, for this
    pool) — normalize to the parsed value either way."""
    if raw is None:
        return None
    return raw if not isinstance(raw, str) else json.loads(raw)


async def fetch_previous_recommendations(state: TaskReportCritiqueState) -> Dict[str, Any]:
    """Look up the previous attempt's recommendations for this (deviation_id, task_index) and
    keep only the ones the investigator accepted, so critique_tasks can check whether this new
    report actually addresses them.

    deviation_id/task_index aren't sent by the backend yet (see
    ds/src/agents/critique/GAPS.md) — until they are, this always returns an empty list, and the
    rest of the graph behaves exactly as it did before this node existed.

    Cross-references two tables because investigation_task_critique_recommendation_history never
    gets a decision written into it (it's logged at generation time, before any decision is
    made), while investigation_task_critique_reports' current row still holds the previous
    attempt's final decisions — the backend calls this endpoint before it overwrites that row for
    the new attempt. Both lists come from the same ordered recommendations list written in the
    same backend call, so they're zipped by index.

    Best-effort throughout: any DB error, missing row, or shape mismatch degrades to an empty
    list rather than failing the critique request.
    """
    if state.deviation_id is None or state.task_index is None:
        return {"previous_recommendations": []}

    try:
        pool = await get_db_pool()
        async with pool.acquire() as conn:
            history_row = await conn.fetchrow(
                """
                SELECT recommendations FROM investigation_task_critique_recommendation_history
                WHERE deviation_id = $1 AND task_index = $2
                ORDER BY created_at DESC LIMIT 1
                """,
                state.deviation_id,
                state.task_index,
            )
            report_row = await conn.fetchrow(
                """
                SELECT recommendations FROM investigation_task_critique_reports
                WHERE deviation_id = $1 AND task_index = $2
                """,
                state.deviation_id,
                state.task_index,
            )
    except Exception:
        logger.warning(
            "fetch_previous_recommendations: DB lookup failed for deviation_id=%s task_index=%s",
            state.deviation_id, state.task_index, exc_info=True,
        )
        return {"previous_recommendations": []}

    if history_row is None or report_row is None:
        return {"previous_recommendations": []}

    history_recs = _parse_json_column(history_row["recommendations"])
    report_recs = _parse_json_column(report_row["recommendations"])
    if not isinstance(history_recs, list) or not isinstance(report_recs, list):
        return {"previous_recommendations": []}
    if len(history_recs) != len(report_recs):
        logger.warning(
            "fetch_previous_recommendations: history/report recommendation count mismatch "
            "(%d vs %d) for deviation_id=%s task_index=%s — skipping carry-forward",
            len(history_recs), len(report_recs), state.deviation_id, state.task_index,
        )
        return {"previous_recommendations": []}

    accepted = [
        description
        for description, report_item in zip(history_recs, report_recs)
        if isinstance(report_item, dict) and report_item.get("decision") == "accepted"
    ]
    logger.info(
        "fetch_previous_recommendations: %d/%d previous recommendations were accepted for "
        "deviation_id=%s task_index=%s",
        len(accepted), len(report_recs), state.deviation_id, state.task_index,
    )
    return {"previous_recommendations": accepted}


async def analyze_images(state: TaskReportCritiqueState) -> Dict[str, Any]:
    """Run vision analysis on each image concurrently via /v1/chat/completions."""
    if not state.image_context_map:
        return {"image_analyses": []}

    llm: LLMClient = await get_llm_client()
    registry: PromptRegistry = get_prompt_registry()
    prompt_template = registry.get("task_report_critique/analyze_image")

    investigation_context = (
        f"Problem statement: {state.report_metadata.get('problem_statement', '')}\n"
        f"Objective: {state.report_metadata.get('objective', '')}"
    )

    async def _analyse_one(img: Dict[str, Any]) -> Dict[str, Any]:
        prompt = prompt_template.format(
            investigation_context=investigation_context,
            context_label=img["context_label"],
            written_claim=img["context_label"],
        )
        try:
            result: ImageAnalysisResult = await llm.get_structured_vision_chat_response(
                user_prompt=prompt,
                image_b64=img["b64"],
                image_media_type=img["media_type"],
                structure=ImageAnalysisResult,
            )
            return {
                "name": img["name"],
                "context_label": img["context_label"],
                "is_evidence_photo": result.is_evidence_photo,
                "observation": result.observation,
                "compliance_concerns": result.compliance_concerns,
                "supports_written_claim": result.supports_written_claim,
            }
        except Exception:
            logger.warning("Vision analysis failed for image %s", img["name"], exc_info=True)
            return {
                "name": img["name"],
                "context_label": img["context_label"],
                "is_evidence_photo": False,
                "observation": "",
                "compliance_concerns": "",
                "supports_written_claim": True,
            }

    analyses = await asyncio.gather(*[_analyse_one(img) for img in state.image_context_map])
    evidence_count = sum(1 for a in analyses if a["is_evidence_photo"])
    logger.info(
        "analyze_images: %d/%d images are evidence photos",
        evidence_count, len(analyses),
    )
    return {"image_analyses": list(analyses)}


def _build_visual_obs_for_task(task: Dict[str, Any], image_analyses: List[Dict[str, Any]]) -> str:
    """Return visual observations that belong to this task, matched via section_labels."""
    section_labels = {lbl.lower() for lbl in task.get("section_labels", [])}
    matched = [
        a for a in image_analyses
        if a.get("is_evidence_photo")
        and a.get("observation")
        and a.get("context_label", "").lower() in section_labels
    ]
    if not matched:
        return "No photographic evidence for this task."
    return "\n".join(
        f"- [{a['context_label']}] {a['observation']}"
        + (f" Compliance concerns: {a['compliance_concerns']}." if a.get("compliance_concerns") else "")
        + (f" Does not support written claim." if not a.get("supports_written_claim") else "")
        for a in matched
    )


async def critique_tasks(state: TaskReportCritiqueState) -> Dict[str, Any]:
    """Apply critique dimensions to every task, passing only that task's visual evidence."""
    llm: LLMClient = await get_llm_client()
    registry: PromptRegistry = get_prompt_registry()

    # Attach per-task visual observations before building the prompt
    tasks_with_visuals = []
    for task in state.extracted_tasks:
        tasks_with_visuals.append({
            **task,
            "visual_observations": _build_visual_obs_for_task(task, state.image_analyses),
        })

    previous_recommendations_text = (
        "\n".join(f"- {rec}" for rec in state.previous_recommendations)
        if state.previous_recommendations
        else "None."
    )

    prompt = registry.get("task_report_critique/critique_task").format(
        problem_statement=state.report_metadata.get("problem_statement", ""),
        objective=state.report_metadata.get("objective", ""),
        tasks_json=json.dumps(tasks_with_visuals, ensure_ascii=False, indent=2),
        previous_recommendations=previous_recommendations_text,
    )

    result: AllTaskCritiquesResult = await llm.get_structured_chat_response(
        user_prompt=prompt,
        structure=AllTaskCritiquesResult,
    )

    for critique in result.task_critiques:
        critique.recommendations.sort(key=lambda r: not r.startswith(UNADDRESSED_MARKER))

    # Cap is per REPORT, not per task: a single uploaded report can be split into several
    # task_critiques entries (one per "Inference:" block found by extract_tasks), but only 5
    # recommendations total should ever surface for the whole report — the prompt already
    # instructs the model to budget across all tasks combined, this is just a code-side backstop
    # in case it doesn't fully honor that. Still-unaddressed carried-forward items take priority,
    # preserving each task's own carried-forward-first order via a stable sort.
    tagged = [
        (task_idx, rec)
        for task_idx, critique in enumerate(result.task_critiques)
        for rec in critique.recommendations
    ]
    tagged.sort(key=lambda pair: not pair[1].startswith(UNADDRESSED_MARKER))
    kept_by_task: Dict[int, List[str]] = defaultdict(list)
    for task_idx, rec in tagged[:5]:
        kept_by_task[task_idx].append(rec)
    for task_idx, critique in enumerate(result.task_critiques):
        critique.recommendations = kept_by_task.get(task_idx, [])

    logger.info("critique_tasks: critiqued %d tasks", len(result.task_critiques))
    return {
        "task_critiques": [c.model_dump() for c in result.task_critiques],
    }


async def format_result(state: TaskReportCritiqueState) -> Dict[str, Any]:
    """Assemble the final API response."""
    from src.agents.critique.graph.schemas import TaskCritiqueDetail

    final = TaskReportCritiqueResponse(
        problem_statement=state.report_metadata.get("problem_statement", ""),
        objective=state.report_metadata.get("objective", ""),
        task_critiques=[TaskCritiqueDetail.model_validate(c) for c in state.task_critiques],
        total_tasks_analyzed=len(state.task_critiques),
    )
    return {"final_result": final.model_dump()}
