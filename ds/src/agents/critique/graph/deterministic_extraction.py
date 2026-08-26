"""Deterministic (non-LLM) task extraction for the Task Report Critique pipeline.

Every real uploaded task-report .docx checked in live testing (2026-08-26, 37 documents across
6 investigation records under ~/Downloads/test/ and ~/Downloads/505542_v2/) has exactly one
"Title of the task:" / "Objective:" / "Inference" marker set — one uploaded document is one
task, always, in current real-world usage. That template is structurally simple enough to parse
without an LLM: the markers are bold, paragraph-leading label paragraphs, and the content
between them is copied verbatim.

This module only ever produces a result when it's confident that grammar holds (exactly one of
each required marker, in order, with non-empty content between them). Anything else — a
different template, a genuinely multi-task document, missing/duplicated markers — returns None,
and the caller (`nodes.py: extract_tasks`) falls back to the existing LLM extraction path
unchanged. No guessing: a wrong silent extraction is worse than a slow correct one in this
regulated pharma-QA context (see ds/src/agents/critique/GAPS.md for prior, unrelated deterministic
parsers in this codebase that needed multiple rounds of live-document correction — this module
deliberately stays narrow in scope to minimize that risk).
"""

from __future__ import annotations

import re
from typing import Any, Dict, List, Optional

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.agents.critique.graph.schemas import ExtractedTask, ExtractionResult


def _is_effectively_bold(run, paragraph) -> bool:
    """run.bold is None whenever boldness comes from the paragraph's style (e.g. Word's built-in
    'Heading N' styles) rather than being set on the run itself — treating None as falsy would
    miss every style-based heading, so fall back to the paragraph's style chain."""
    if run.bold is not None:
        return run.bold
    style = paragraph.style
    while style is not None:
        if style.font.bold is not None:
            return style.font.bold
        style = style.base_style
    return False


# Each marker must be the paragraph's ENTIRE text (anchored start+end, after stripping) AND its
# leading run must be bold — this is what rejects "Inference" appearing mid-sentence inside a
# "[Source: ..., Section 6 – Method / Inference, p.18]" citation (not bold-leading, not the whole
# paragraph) as well as a different template's inline "Inference: <body text on the same line>"
# style (fails the whole-paragraph match, correctly triggering the None/LLM-fallback path instead
# of a wrong guess).
_MARKER_PATTERNS: List[tuple[str, re.Pattern]] = [
    ("PROBLEM_STATEMENT", re.compile(r'^\**\s*Problem\s+statement\s*:?\**\s*$', re.I)),
    ("INVESTIGATION_TASKS", re.compile(r'^\**\s*Investigation\s+tasks\s*:?\**\s*$', re.I)),
    ("TITLE", re.compile(r'^\**\s*Title\s+of\s+the\s+task\s*:?\**\s*$', re.I)),
    ("OBJECTIVE", re.compile(r'^\**\s*Objective\s*:?\**\s*$', re.I)),
    ("TASK_DETAILS", re.compile(r'^\**\s*Task\s+details\s*:?\**\s*$', re.I)),
    ("FINDINGS", re.compile(r'^\**\s*Findings\s*:?\**\s*$', re.I)),
    ("INFERENCE", re.compile(r'^\**\s*Inference\s*:?\**\s*$', re.I)),
]
# Markers that are pure structural labels within the findings body — recognized (as boundaries
# for objective's own span) but excluded from the extracted findings text itself, same as every
# other marker is excluded from its own following content.
_FINDINGS_BOUNDARY_MARKERS = ("TASK_DETAILS", "FINDINGS")


def _classify_marker(paragraph: Paragraph) -> Optional[str]:
    text = paragraph.text.strip()
    if not text:
        return None
    runs = [r for r in paragraph.runs if r.text and r.text.strip()]
    if not runs or not _is_effectively_bold(runs[0], paragraph):
        return None
    for kind, pattern in _MARKER_PATTERNS:
        if pattern.match(text):
            return kind
    return None


def count_true_inference_markers(file_path: str) -> int:
    """Bold-aware, structural count of genuine Inference markers, for callers (the LLM
    extraction path) that need to tell the model how many tasks to expect. Replaces a
    markdown-text regex count that matched any paragraph starting with "Inference:" regardless
    of formatting — which over-counts on documents that use a non-bold "Inference: <sentence>"
    as an inline sub-conclusion label within findings prose, a distinct pattern from the
    mid-citation false positive fixed earlier. Confirmed live, 2026-08-26, on
    "503442_4 MEASUREMENT...docx": true count 1, text-based count 5 — the inflated count caused
    the LLM to fabricate 5 tasks, inventing inference text for 4 of them, where the source
    document's single real Inference section is actually blank (a genuine content gap, not an
    extraction bug — the LLM fabricated plausible-sounding text to paper over it)."""
    doc = Document(file_path)
    return sum(1 for p in doc.paragraphs if _classify_marker(p) == "INFERENCE")


def _table_to_text(table: Table) -> str:
    rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
    return "\n".join(r for r in rows if r.strip())


# "Figure N: <caption>" paragraphs are pure image references with no analytical content — the
# actual photo content is captured separately by `analyze_images`/`visual_observations`, so these
# add nothing but bulk. The LLM extraction path consistently omits them from `findings`
# (confirmed live, 2026-08-26: on 500954_1, findings similarity vs. the LLM baseline was 0.42
# with these lines included, 1.00 once stripped — every other word matched) — excluded here to
# match that established behavior rather than diverge from it for no benefit.
_FIGURE_CAPTION_RE = re.compile(r'^Figure\s+\d+\s*:', re.I)


def _span_text(items: List[Any], start: int, end: int, exclude: set[int]) -> str:
    """Concatenate paragraph/table text for items[start:end], skipping marker paragraphs and
    figure-caption paragraphs."""
    parts: List[str] = []
    for idx in range(start, end):
        if idx in exclude:
            continue
        item = items[idx]
        if isinstance(item, Paragraph):
            t = item.text.strip()
            if t and not _FIGURE_CAPTION_RE.match(t):
                parts.append(t)
        elif isinstance(item, Table):
            t = _table_to_text(item)
            if t:
                parts.append(t)
    return "\n".join(parts).strip()


def _bold_headings_in_span(items: List[Any], start: int, end: int, exclude: set[int]) -> List[str]:
    """Bold sub-heading paragraphs and table first-row/first-cell text within the span — the same
    signal `_extract_images_from_docx` (nodes.py) uses to label each image's context, so labels
    collected here should naturally line up with real images' context_label for the fuzzy matching
    `_labels_match` already does downstream."""
    labels: List[str] = []
    for idx in range(start, end):
        if idx in exclude:
            continue
        item = items[idx]
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            runs = [r for r in item.runs if r.text and r.text.strip()]
            # A real section heading in this template is a short phrase; a bold full paragraph
            # (e.g. a "PARTIAL EVIDENCE GAP - ..." callout) is body content, not a heading — cap
            # length so callouts don't leak into section_labels (confirmed live, 2026-08-26).
            if (
                runs
                and _is_effectively_bold(runs[0], item)
                and _classify_marker(item) is None
                and len(text) <= 120
            ):
                labels.append(text)
        elif isinstance(item, Table) and item.rows:
            header = item.rows[0].cells[0].text.strip()
            if header and len(header) <= 120:
                labels.append(header)
    # Dedupe, preserve order.
    seen: set[str] = set()
    deduped = []
    for label in labels:
        if label not in seen:
            seen.add(label)
            deduped.append(label)
    return deduped


def extract_task_deterministic(file_path: str) -> Optional[ExtractionResult]:
    """Try to extract a single task deterministically. Returns None (defer to LLM extraction)
    unless exactly one TITLE/OBJECTIVE/INFERENCE marker is found, in order, with non-empty
    content in every required span."""
    doc = Document(file_path)
    items = list(doc.iter_inner_content())

    markers: Dict[str, List[int]] = {}
    for idx, item in enumerate(items):
        if isinstance(item, Paragraph):
            kind = _classify_marker(item)
            if kind:
                markers.setdefault(kind, []).append(idx)

    for required in ("TITLE", "OBJECTIVE", "INFERENCE"):
        if len(markers.get(required, [])) != 1:
            return None

    title_idx = markers["TITLE"][0]
    objective_idx = markers["OBJECTIVE"][0]
    inference_idx = markers["INFERENCE"][0]
    if not (title_idx < objective_idx < inference_idx):
        return None

    problem_statement_idx = markers.get("PROBLEM_STATEMENT", [None])[0]
    investigation_tasks_idx = markers.get("INVESTIGATION_TASKS", [None])[0]
    if problem_statement_idx is not None and problem_statement_idx >= title_idx:
        return None
    if investigation_tasks_idx is not None and not (
        (problem_statement_idx is None or problem_statement_idx < investigation_tasks_idx)
        and investigation_tasks_idx < title_idx
    ):
        return None

    # Objective's span ends at the first of TASK_DETAILS/FINDINGS/INFERENCE that follows it.
    boundary_candidates = [
        idx
        for kind in (*_FINDINGS_BOUNDARY_MARKERS, "INFERENCE")
        for idx in markers.get(kind, [])
        if idx > objective_idx
    ]
    if not boundary_candidates:
        return None
    objective_end_idx = min(boundary_candidates)

    all_marker_indices = {idx for idxs in markers.values() for idx in idxs}

    title_text = ""
    for idx in range(title_idx + 1, objective_idx):
        if idx in all_marker_indices:
            break
        item = items[idx]
        if isinstance(item, Paragraph) and item.text.strip():
            title_text = item.text.strip()
            break

    problem_statement = ""
    if problem_statement_idx is not None:
        ps_end = investigation_tasks_idx if investigation_tasks_idx is not None else title_idx
        problem_statement = _span_text(items, problem_statement_idx + 1, ps_end, all_marker_indices)

    doc_objective = ""
    if investigation_tasks_idx is not None:
        doc_objective = _span_text(items, investigation_tasks_idx + 1, title_idx, all_marker_indices)

    # Findings start at the FINDINGS marker specifically, not TASK_DETAILS — the RCI Plan's
    # "Task details:" bullets restate the task's original scope/instructions (copied for
    # reference), not actual findings; the LLM baseline consistently excludes them from
    # `findings` (confirmed live, 2026-08-26, on 505542_2), so drop that span here too rather
    # than diverging from established behavior.
    findings_markers_after = sorted(idx for idx in markers.get("FINDINGS", []) if idx > objective_idx)
    findings_start_idx = findings_markers_after[0] + 1 if findings_markers_after else objective_end_idx

    # A trailing "Prepared by / Name / Sign/Date" signature table is a document-footer artifact
    # of this template, not part of the inference — confirmed present (in this exact form) at
    # the end of every sample document checked. Cap the inference span before it.
    inference_end_idx = len(items)
    for idx in range(inference_idx + 1, len(items)):
        item = items[idx]
        if isinstance(item, Table) and item.rows and item.rows[0].cells:
            if "prepared by" in item.rows[0].cells[0].text.strip().lower():
                inference_end_idx = idx
                break

    objective = _span_text(items, objective_idx + 1, objective_end_idx, all_marker_indices)
    findings = _span_text(items, findings_start_idx, inference_idx, all_marker_indices)
    inference = _span_text(items, inference_idx + 1, inference_end_idx, all_marker_indices)
    section_labels = _bold_headings_in_span(items, findings_start_idx, inference_idx, all_marker_indices)

    if not all([title_text, objective, findings, inference]):
        return None

    task = ExtractedTask(
        task_number=1,
        title=title_text,
        objective=objective,
        findings=findings,
        inference=inference,
        section_labels=section_labels,
    )
    return ExtractionResult(problem_statement=problem_statement, objective=doc_objective, tasks=[task])


# ── Legacy template (fallback tier 2) ───────────────────────────────────────
#
# An older task-report template, confirmed live (2026-08-26) against fixtures in
# InvestigationAi_DS/data/critique/: no "Title of the task:" marker (the title is just the
# first bold paragraph after "Investigation tasks"), and Objective:/Findings:/Inference: are
# inconsistently bold — sometimes bold, sometimes plain, varying paragraph to paragraph even
# within the same document. Critically, the marker text often shares its OWN paragraph with the
# field's content (e.g. "Objective: To identify if..." on one line), unlike the current
# template where a marker is always alone on its own paragraph. Confirmed single-task only
# (per user, 2026-08-26) — some sibling fixtures in that folder are actually full combined RCI
# reports (task sections followed by Root Cause/Impact Assessment/CAPA content) rather than
# single task-report uploads, so this deliberately does NOT attempt multi-task segmentation;
# a document like that will fail the single-marker-count check below and correctly defer to
# the LLM path rather than risk pulling CAPA/root-cause content into `findings`.

_LEGACY_MARKER_PATTERNS: List[tuple[str, re.Pattern]] = [
    ("PROBLEM_STATEMENT", re.compile(r'^\**\s*(?:\d+\.?\s*)?Problem\s+statement\s*:?\**\s*', re.I)),
    ("INVESTIGATION_TASKS", re.compile(r'^\**\s*(?:\d+\.?\s*)?Investigation\s+tasks?\s*:?\**\s*', re.I)),
    ("OBJECTIVE", re.compile(r'^\**\s*Objective\s*:?\**\s*', re.I)),
    ("FINDINGS", re.compile(r'^\**\s*Findings\s*:?\**\s*', re.I)),
    ("INFERENCE", re.compile(r'^\**\s*Inference\s*:?\**\s*', re.I)),
]


def _classify_legacy_marker(paragraph: Paragraph) -> Optional[tuple[str, str]]:
    """Prefix match at the start of the paragraph, regardless of bold formatting. Returns
    (marker_kind, inline_trailing_text) so content sharing the marker's own paragraph — common
    in this template — isn't lost."""
    text = paragraph.text.strip()
    if not text:
        return None
    for kind, pattern in _LEGACY_MARKER_PATTERNS:
        m = pattern.match(text)
        if m:
            return kind, text[m.end():].strip()
    return None


def extract_task_deterministic_legacy(file_path: str) -> Optional[ExtractionResult]:
    """Second-tier deterministic extraction for the older template. Only attempted after
    `extract_task_deterministic` returns None. Same never-guess contract: returns None on
    anything ambiguous (no anchor for the front matter, no title candidate, multiple
    Objective/Inference markers, or any empty required field) rather than risk a wrong
    extraction — the caller falls through to the LLM path."""
    doc = Document(file_path)
    items = list(doc.iter_inner_content())

    markers: Dict[str, List[int]] = {}
    inline: Dict[int, str] = {}
    for idx, item in enumerate(items):
        if isinstance(item, Paragraph):
            classified = _classify_legacy_marker(item)
            if classified:
                kind, trailing = classified
                markers.setdefault(kind, []).append(idx)
                inline[idx] = trailing

    for required in ("OBJECTIVE", "INFERENCE"):
        if len(markers.get(required, [])) != 1:
            return None

    objective_idx = markers["OBJECTIVE"][0]
    inference_idx = markers["INFERENCE"][0]
    if objective_idx >= inference_idx:
        return None

    problem_statement_idx = markers.get("PROBLEM_STATEMENT", [None])[0]
    investigation_tasks_idx = markers.get("INVESTIGATION_TASKS", [None])[0]
    if problem_statement_idx is not None and problem_statement_idx >= objective_idx:
        return None
    if investigation_tasks_idx is not None and investigation_tasks_idx >= objective_idx:
        return None
    # Require an anchor for the front matter — scanning for a title from paragraph 0 in a
    # document with neither marker is unbounded and too risky to trust.
    if problem_statement_idx is None and investigation_tasks_idx is None:
        return None

    all_marker_indices = {idx for idxs in markers.values() for idx in idxs}

    # Title: the first bold paragraph after the front matter that isn't itself a marker — this
    # template has no dedicated title label.
    title_search_start = (
        investigation_tasks_idx + 1 if investigation_tasks_idx is not None
        else problem_statement_idx + 1
    )
    title_text = ""
    for idx in range(title_search_start, objective_idx):
        if idx in all_marker_indices:
            continue
        item = items[idx]
        if not isinstance(item, Paragraph):
            continue
        text = item.text.strip()
        if not text:
            continue
        runs = [r for r in item.runs if r.text and r.text.strip()]
        if runs and _is_effectively_bold(runs[0], item):
            title_text = text
            break

    findings_markers_after = sorted(
        idx for idx in markers.get("FINDINGS", []) if objective_idx < idx < inference_idx
    )
    findings_marker_idx = findings_markers_after[0] if findings_markers_after else None
    objective_end_idx = findings_marker_idx if findings_marker_idx is not None else inference_idx
    findings_body_start = (findings_marker_idx + 1) if findings_marker_idx is not None else objective_end_idx

    problem_statement = ""
    if problem_statement_idx is not None:
        ps_end = investigation_tasks_idx if investigation_tasks_idx is not None else title_search_start
        problem_statement = "\n".join(filter(None, [
            inline.get(problem_statement_idx, ""),
            _span_text(items, problem_statement_idx + 1, ps_end, all_marker_indices),
        ])).strip()

    objective = "\n".join(filter(None, [
        inline.get(objective_idx, ""),
        _span_text(items, objective_idx + 1, objective_end_idx, all_marker_indices),
    ])).strip()

    findings = "\n".join(filter(None, [
        inline.get(findings_marker_idx, "") if findings_marker_idx is not None else "",
        _span_text(items, findings_body_start, inference_idx, all_marker_indices),
    ])).strip()

    inference = "\n".join(filter(None, [
        inline.get(inference_idx, ""),
        _span_text(items, inference_idx + 1, len(items), all_marker_indices),
    ])).strip()

    section_labels = _bold_headings_in_span(items, findings_body_start, inference_idx, all_marker_indices)

    if not all([title_text, objective, findings, inference]):
        return None

    task = ExtractedTask(
        task_number=1,
        title=title_text,
        objective=objective,
        findings=findings,
        inference=inference,
        section_labels=section_labels,
    )
    return ExtractionResult(problem_statement=problem_statement, objective="", tasks=[task])
