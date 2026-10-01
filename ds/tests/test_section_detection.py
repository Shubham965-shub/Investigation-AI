"""Tests for section_detection's pure helpers and local text-extraction fallback.
docx_to_text uses python-docx against a real generated .docx file (no LLM/DB
involved) — a genuine parsing-logic test, not a mock of the library."""

from docx import Document

from src.agents.scoring.api.schemas import SectionDetectionLLMOutput
from src.agents.scoring.services.section_detection import (
    docx_to_text,
    normalise_event_type,
    present_sections,
)


class TestNormaliseEventType:
    def test_none_or_empty_returns_none(self):
        assert normalise_event_type(None) is None
        assert normalise_event_type("") is None

    def test_known_aliases_canonicalize(self):
        assert normalise_event_type("oos") == "OOS"
        assert normalise_event_type("Out Of Specification") == "OOS"
        assert normalise_event_type("out of trend") == "OOT"
        assert normalise_event_type("complaint") == "Market Complaint"
        assert normalise_event_type("Market Complaint") == "Market Complaint"

    def test_unknown_value_is_stripped_but_passed_through(self):
        assert normalise_event_type("  Some Other Type  ") == "Some Other Type"


class TestPresentSections:
    def test_only_non_blank_sections_are_included(self):
        detection = SectionDetectionLLMOutput(
            task_report_text="  has content  ",
            rc_text="",
            impact_text="   ",
            capa_text="capa content",
        )
        result = present_sections(detection)
        assert result == {"task_report": "has content", "capa": "capa content"}

    def test_all_blank_returns_empty_dict(self):
        detection = SectionDetectionLLMOutput()
        assert present_sections(detection) == {}


class TestDocxToText(object):
    def test_extracts_paragraphs_and_table_rows(self, tmp_path):
        doc = Document()
        doc.add_paragraph("Task Report Execution")
        doc.add_paragraph("")  # blank paragraph should be skipped
        doc.add_paragraph("Findings are supported by objective evidence.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Checkpoint"
        table.cell(0, 1).text = "Verdict"
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "Yes"
        file_path = tmp_path / "report.docx"
        doc.save(str(file_path))

        text = docx_to_text(file_path)

        assert "Task Report Execution" in text
        assert "Findings are supported by objective evidence." in text
        assert "Checkpoint | Verdict" in text
        assert "1 | Yes" in text

    def test_empty_document_returns_empty_string(self, tmp_path):
        doc = Document()
        file_path = tmp_path / "empty.docx"
        doc.save(str(file_path))
        assert docx_to_text(file_path) == ""
