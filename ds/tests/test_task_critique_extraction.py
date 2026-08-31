"""
Regression tests for the Task Report Critique feature's deterministic (non-LLM)
extraction path (src/agents/critique/graph/deterministic_extraction.py) and the
pure markdown/image helpers it shares with src/agents/critique/graph/nodes.py.

GAPS.md's 2026-08-26 entries document several real bugs found live: an
un-anchored "Inference:" regex overcounting real task markers (causing the LLM
to fabricate extra tasks), a second false-positive pattern from non-bold inline
"Inference: <sentence>" sub-conclusions, and style-based bold (Word's built-in
Heading styles, where `run.bold` is `None`) needing a fallback through the
paragraph's style chain. None of these had a dedicated test before this file.
"""

import io
import tempfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from src.agents.app import create_app
from src.agents.critique.api.services.relevance_validation import DocumentRelevanceResult
from src.agents.critique.graph.deterministic_extraction import (
    _classify_marker,
    _is_effectively_bold,
    count_true_inference_markers,
    extract_task_deterministic,
)
from src.agents.critique.graph.nodes import _resize_and_encode_image, _segment_sections, _task_body
from src.agents.critique.graph.schemas import AllTaskCritiquesResult, TaskCritiqueDetail
from src.llm.client import LLMClient
from src.utils import deps


# ---------------------------------------------------------------------------
# _is_effectively_bold / _classify_marker
# ---------------------------------------------------------------------------


def _bold_paragraph(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    return p


def _plain_paragraph(doc, text):
    return doc.add_paragraph(text)


def test_classify_marker_requires_bold_leading_run():
    doc = Document()
    plain = _plain_paragraph(doc, "Inference")
    bold = _bold_paragraph(doc, "Inference")
    assert _classify_marker(plain) is None
    assert _classify_marker(bold) == "INFERENCE"


def test_classify_marker_rejects_mid_citation_occurrence():
    """The exact false-positive documented in GAPS.md: 'Inference' appearing
    inside a citation string is not bold-leading and not the whole paragraph."""
    doc = Document()
    p = doc.add_paragraph("[Source: RCI Plan, Section 6 - Method / Milling / Inference, p.18]")
    assert _classify_marker(p) is None


def test_classify_marker_rejects_inline_marker_with_trailing_content():
    """A different template's 'Inference: <body text on the same line>' style
    fails the whole-paragraph match, correctly deferring to the LLM path
    instead of a wrong guess."""
    doc = Document()
    p = _bold_paragraph(doc, "Inference: the root cause is a torque wrench calibration lapse.")
    assert _classify_marker(p) is None


def test_is_effectively_bold_falls_back_to_style_chain():
    """run.bold is None when boldness comes from the paragraph's style (e.g.
    Word's built-in Heading styles) rather than the run itself — treating None
    as falsy would miss every style-based heading."""
    doc = Document()
    p = doc.add_paragraph("Inference")
    p.style = doc.styles["Heading 2"]  # built-in heading styles are bold by default
    run = p.runs[0]
    assert run.bold is None
    assert _is_effectively_bold(run, p) is True


def test_classify_marker_matches_via_style_based_bold_heading():
    doc = Document()
    p = doc.add_paragraph("Inference")
    p.style = doc.styles["Heading 2"]
    assert _classify_marker(p) == "INFERENCE"


# ---------------------------------------------------------------------------
# count_true_inference_markers
# ---------------------------------------------------------------------------


def _save(doc: Document) -> Path:
    tmp = Path(tempfile.mkstemp(suffix=".docx")[1])
    doc.save(tmp)
    return tmp


def test_count_true_inference_markers_ignores_non_bold_inline_sub_conclusions():
    """Reproduces the second bug (2026-08-26): a document with one real bold
    'Inference' marker plus several non-bold inline 'Inference: <sentence>'
    sub-conclusions scattered through the findings prose. The old markdown-text
    regex count was inflated by these; the bold-aware structural count must
    return exactly 1."""
    doc = Document()
    doc.add_paragraph("Title of the task: Verify dispensing area SOP compliance.")
    doc.add_paragraph("Objective: Determine adherence to the dispensing SOP.")
    doc.add_paragraph("Findings: Line clearance was verified across three shifts.")
    doc.add_paragraph("Inference: partial compliance noted in shift 2 records.")  # non-bold inline
    doc.add_paragraph("Inference: no deviation found in shift 3 records.")  # non-bold inline
    _bold_paragraph(doc, "Inference")  # the one true structural marker
    doc.add_paragraph("No deviation from the approved SOP was found overall.")

    path = _save(doc)
    try:
        assert count_true_inference_markers(str(path)) == 1
    finally:
        path.unlink()


def test_count_true_inference_markers_counts_zero_when_none_bold():
    doc = Document()
    doc.add_paragraph("Inference: informal note only, never a real structural marker.")
    path = _save(doc)
    try:
        assert count_true_inference_markers(str(path)) == 0
    finally:
        path.unlink()


# ---------------------------------------------------------------------------
# extract_task_deterministic — end-to-end grammar match / decline
# ---------------------------------------------------------------------------


def _well_formed_single_task_doc() -> Document:
    doc = Document()
    _bold_paragraph(doc, "Problem statement:")
    doc.add_paragraph("Foreign object found embedded in a tablet during in-process inspection.")
    _bold_paragraph(doc, "Investigation tasks")
    doc.add_paragraph("Verify the dispensing area's material handling controls.")
    _bold_paragraph(doc, "Title of the task:")
    doc.add_paragraph("VERIFICATION OF DISPENSING AREA SOP")
    _bold_paragraph(doc, "Objective:")
    doc.add_paragraph("Determine whether the dispensing SOP was followed for the affected batch.")
    _bold_paragraph(doc, "Findings:")
    doc.add_paragraph("Line clearance records for all three shifts were reviewed and found complete.")
    _bold_paragraph(doc, "Inference")
    doc.add_paragraph("No SOP deviation was identified in the dispensing area for this batch.")
    return doc


def test_extract_task_deterministic_matches_well_formed_single_task_document():
    path = _save(_well_formed_single_task_doc())
    try:
        result = extract_task_deterministic(str(path))
    finally:
        path.unlink()

    assert result is not None
    assert len(result.tasks) == 1
    task = result.tasks[0]
    assert task.title == "VERIFICATION OF DISPENSING AREA SOP"
    assert "line clearance records" in task.findings.lower()
    assert "no sop deviation" in task.inference.lower()
    assert "foreign object" in result.problem_statement.lower()


def test_extract_task_deterministic_declines_document_with_two_inference_markers():
    """A genuinely multi-task (or otherwise out-of-grammar) document must
    return None rather than guess — the caller then falls back to the LLM
    extraction path unchanged."""
    doc = _well_formed_single_task_doc()
    _bold_paragraph(doc, "Inference")  # a second structural marker
    doc.add_paragraph("A second, unexpected inference paragraph.")

    path = _save(doc)
    try:
        result = extract_task_deterministic(str(path))
    finally:
        path.unlink()

    assert result is None


def test_extract_task_deterministic_declines_when_findings_span_is_empty():
    doc = Document()
    _bold_paragraph(doc, "Title of the task:")
    doc.add_paragraph("VERIFICATION OF DISPENSING AREA SOP")
    _bold_paragraph(doc, "Objective:")
    doc.add_paragraph("Determine whether the dispensing SOP was followed.")
    _bold_paragraph(doc, "Findings:")
    # no findings content at all before Inference
    _bold_paragraph(doc, "Inference")
    doc.add_paragraph("No deviation identified.")

    path = _save(doc)
    try:
        result = extract_task_deterministic(str(path))
    finally:
        path.unlink()

    assert result is None


# ---------------------------------------------------------------------------
# nodes.py — _task_body / _segment_sections (pure markdown text functions)
# ---------------------------------------------------------------------------


def test_task_body_truncates_at_identified_root_cause_heading():
    markdown = (
        "1. Material\nObjective: check vendor lot.\nInference: no defect found.\n\n"
        "Identified Root Cause: Deficient incoming inspection procedure.\n"
        "CAPA Proposal: Revise SOP.\n"
    )
    body = _task_body(markdown)
    assert "Identified Root Cause" not in body
    assert "no defect found" in body


def test_task_body_returns_full_text_when_no_end_heading_present():
    markdown = "1. Material\nObjective: check vendor lot.\nInference: no defect found.\n"
    assert _task_body(markdown) == markdown


def test_task_body_does_not_truncate_on_causal_factor_body_text():
    """'Causal Factor' requires a trailing colon to be treated as an end
    heading — must not match ordinary body prose like 'causal factors can be
    ruled out.'"""
    markdown = "1. Material\nInference: causal factors can be ruled out for this task.\n"
    assert _task_body(markdown) == markdown


def test_segment_sections_splits_on_top_level_numbered_headings_only():
    markdown = (
        "1. Material\nObjective: check vendor lot.\nInference: no defect found.\n\n"
        "1.1 Sub-point\nThis should not start a new top-level section.\n\n"
        "2. Method\nObjective: verify SOP.\nInference: SOP was followed correctly.\n"
    )
    segments = _segment_sections(markdown)
    assert len(segments) == 2
    assert segments[0][0].startswith("1.")
    assert "no defect found" in segments[0][1]
    assert segments[1][0].startswith("2.")
    assert "sop was followed correctly" in segments[1][1].lower()


def test_segment_sections_excludes_sections_without_inference_marker():
    markdown = (
        "1. Material\nObjective: check vendor lot. No inference recorded yet.\n\n"
        "2. Method\nObjective: verify SOP.\nInference: SOP was followed correctly.\n"
    )
    segments = _segment_sections(markdown)
    assert len(segments) == 1
    assert segments[0][0].startswith("2.")


def test_segment_sections_no_numbered_headings_returns_whole_body_as_one_segment():
    markdown = "Just a single unstructured findings block with no numbered headings."
    assert _segment_sections(markdown) == [("", markdown)]


# ---------------------------------------------------------------------------
# nodes.py — _resize_and_encode_image (pure image transform)
# ---------------------------------------------------------------------------


def _png_bytes(width: int, height: int) -> bytes:
    from io import BytesIO

    from PIL import Image

    buf = BytesIO()
    Image.new("RGB", (width, height), color=(200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def test_resize_and_encode_image_leaves_small_image_untouched():
    import base64

    blob = _png_bytes(400, 300)
    encoded, content_type = _resize_and_encode_image(blob, "image/png")
    assert content_type == "image/png"
    assert base64.b64decode(encoded) == blob


def test_resize_and_encode_image_downscales_oversized_image_to_jpeg():
    from io import BytesIO

    from PIL import Image

    blob = _png_bytes(3000, 1500)
    encoded, content_type = _resize_and_encode_image(blob, "image/png")
    assert content_type == "image/jpeg"

    import base64

    decoded = Image.open(BytesIO(base64.b64decode(encoded)))
    assert max(decoded.size) <= 2048


def test_resize_and_encode_image_invalid_bytes_falls_back_to_original():
    import base64

    blob = b"not a real image"
    encoded, content_type = _resize_and_encode_image(blob, "image/png")
    assert content_type == "image/png"
    assert base64.b64decode(encoded) == blob


# ---------------------------------------------------------------------------
# Endpoint-level test — POST /critique/analyse-task-report, real 7-node graph
#
# Uses a well-formed single-task .docx so extract_tasks succeeds
# deterministically (confirmed above: zero LLM calls for that node) and has
# no embedded images so analyze_images also short-circuits with no LLM call
# (confirmed via nodes.py:529-530). validate_relevance and critique_tasks
# always call the LLM regardless — those two are mocked via LLMClient's
# class methods, the same class-level monkeypatch pattern used in
# test_capa_depth_effectiveness.py. deviation_id/task_index are left unset so
# fetch_previous_recommendations short-circuits with no DB call at all.
# ---------------------------------------------------------------------------


client = TestClient(create_app())


class _FakePromptRegistry:
    class _Template:
        def format(self, **kwargs):
            return "formatted prompt"

    def get(self, name):
        return self._Template()


@pytest.fixture(autouse=True)
def _reset_deps():
    yield
    deps._llm = None
    deps._prompt_registry = None
    deps._pool = None


def test_analyse_task_report_end_to_end_via_deterministic_extraction(monkeypatch):
    """UT-011/012/013-style scenario: a real single-task .docx flows through
    the full graph (parse_document -> {validate_relevance, extract_tasks,
    fetch_previous_recommendations} -> analyze_images -> critique_tasks ->
    format_result) and returns a critique keyed off the deterministically
    extracted task, not an LLM-fabricated one."""

    async def fake_get_structured_response(self, user_prompt, structure, system_prompt=None, temperature=None):
        assert structure is DocumentRelevanceResult
        return DocumentRelevanceResult(is_relevant=True, reason="Genuine pharma investigation task report.")

    async def fake_get_structured_chat_response(self, user_prompt, structure, system_prompt=None):
        assert structure is AllTaskCritiquesResult
        return AllTaskCritiquesResult(
            task_critiques=[
                TaskCritiqueDetail(
                    task_number=1,
                    title="VERIFICATION OF DISPENSING AREA SOP",
                    recommendations=["Attach the line-clearance checklist as supporting evidence."],
                    strengths="Objective and findings are clearly documented with a traceable conclusion.",
                )
            ]
        )

    monkeypatch.setattr("src.llm.client.LLMClient.get_structured_response", fake_get_structured_response)
    monkeypatch.setattr("src.llm.client.LLMClient.get_structured_chat_response", fake_get_structured_chat_response)
    deps.set_llm(LLMClient())
    deps.set_prompt_registry(_FakePromptRegistry())

    doc = _well_formed_single_task_doc()
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    response = client.post(
        "/critique/analyse-task-report",
        data={"problem_statement": "Foreign object found embedded in a tablet.", "event_type": "Deviation"},
        files={
            "file": (
                "task_report.docx",
                buf.read(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["total_tasks_analyzed"] == 1
    assert data["task_critiques"][0]["title"] == "VERIFICATION OF DISPENSING AREA SOP"
    assert data["task_critiques"][0]["strengths"]
    assert "line-clearance checklist" in data["task_critiques"][0]["recommendations"][0]
    # task_evidence surfaces the deterministically-extracted content verbatim —
    # not re-derived by the critique LLM (see graph/schemas.py's own docstring).
    assert data["task_evidence"][0]["inference"]


def test_analyse_task_report_rejects_non_docx_file():
    response = client.post(
        "/critique/analyse-task-report",
        data={"problem_statement": "Test", "event_type": "Deviation"},
        files={"file": ("report.txt", b"not a docx", "text/plain")},
    )
    assert response.status_code == 415


def test_analyse_task_report_returns_422_when_document_is_irrelevant(monkeypatch):
    async def fake_get_structured_response(self, user_prompt, structure, system_prompt=None, temperature=None):
        assert structure is DocumentRelevanceResult
        return DocumentRelevanceResult(is_relevant=False, reason="Document is a purchase invoice, not an investigation report.")

    monkeypatch.setattr("src.llm.client.LLMClient.get_structured_response", fake_get_structured_response)
    deps.set_llm(LLMClient())
    deps.set_prompt_registry(_FakePromptRegistry())

    doc = _well_formed_single_task_doc()
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    response = client.post(
        "/critique/analyse-task-report",
        data={"problem_statement": "Foreign object found embedded in a tablet.", "event_type": "Deviation"},
        files={
            "file": (
                "task_report.docx",
                buf.read(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 422
