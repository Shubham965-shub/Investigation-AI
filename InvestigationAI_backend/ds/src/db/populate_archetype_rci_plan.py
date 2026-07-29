import os
import re

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from archetype_models import Base, ArchetypeType, Archetype, RciPlan, RciPlanSection, RciPlanTask

DATABASE_URL = "postgresql://postgres:postgres@localhost:5432/app"

EXCEL_PATH = os.path.join(
    os.path.dirname(__file__),
    "RCI_Plan_examples__4__corrected.xlsx",
)

engine = create_engine(DATABASE_URL)
Session = sessionmaker(bind=engine)
session = Session()


def split_title(raw: str) -> tuple[str, str | None]:
    """Split a 'Task(s) & objective' cell into (title, correlation).

    Cells are authored as 'Title (explanation)' or 'Title\\n(explanation)',
    with the explanation itself sometimes containing nested parentheses
    (e.g. '...(knowledge gap, perception gap, situational factors)'). Only
    the outermost paren pair delimits the correlation, so we split on the
    *first* '(' and strip exactly one leading/trailing paren rather than
    matching parens greedily.
    """
    text = raw.strip()
    idx = text.find("(")
    if idx == -1:
        return re.sub(r"\s+", " ", text).strip(), None
    title = re.sub(r"\s+", " ", text[:idx]).strip()
    correlation = text[idx:]
    if correlation.startswith("(") and correlation.endswith(")"):
        correlation = correlation[1:-1]
    else:
        correlation = correlation[1:]
    correlation = re.sub(r"\s+", " ", correlation).strip()
    return title, correlation or None


def load_rci_plans_from_excel(path: str) -> dict[str, list[tuple[str, str | None, list[str]]]]:
    """Return {archetype_name: [(section_title, correlation, [task, ...]), ...]}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    plans: dict[str, list[tuple[str, str | None, list[str]]]] = {}
    for sheet_name in wb.sheetnames:
        if sheet_name == "Index":
            continue
        ws = wb[sheet_name]
        sections: list[tuple[str, str | None, list[str]]] = []
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
            _, objective, detail = (row + (None, None, None))[:3]
            if objective:
                title, correlation = split_title(str(objective))
                sections.append((title, correlation, []))
            if detail and sections:
                sections[-1][2].append(str(detail).strip())
        plans[sheet_name.strip()] = sections
    return plans


# 1. Ensure archetype_type exists
archetype_type = session.query(ArchetypeType).filter_by(name="RCI Plan").first()
if not archetype_type:
    archetype_type = ArchetypeType(name="RCI Plan")
    session.add(archetype_type)
    session.commit()

# 2. Delete existing archetypes (and cascaded plans/sections/tasks) for this type
existing = session.query(Archetype).filter_by(archetype_type_id=archetype_type.id).all()
for arch in existing:
    plan = session.query(RciPlan).filter_by(archetype_id=arch.id).first()
    if plan:
        session.delete(plan)  # cascades to sections -> tasks
    session.delete(arch)
session.commit()
print(f"Removed {len(existing)} existing archetype(s) for 'RCI Plan'.")

# 3. Load new data from Excel
plans = load_rci_plans_from_excel(EXCEL_PATH)

# 4. Insert new archetypes with their RCI plan / sections / tasks
for name, sections in plans.items():
    archetype = Archetype(name=name, definition=None, archetype_type_id=archetype_type.id)
    session.add(archetype)
    session.flush()

    rci_plan = RciPlan(archetype_id=archetype.id)
    session.add(rci_plan)
    session.flush()

    for title, correlation, tasks in sections:
        section = RciPlanSection(rci_plan_id=rci_plan.id, title=title, correlation=correlation)
        session.add(section)
        session.flush()
        for task_desc in tasks:
            session.add(RciPlanTask(rci_plan_section_id=section.id, description=task_desc))

session.commit()
total_sections = sum(len(s) for s in plans.values())
total_tasks = sum(len(t) for _, _, t in (sec for s in plans.values() for sec in s))
print(
    f"Populated {len(plans)} archetype(s) with {total_sections} section(s) and "
    f"{total_tasks} task(s) from '{os.path.basename(EXCEL_PATH)}'."
)
